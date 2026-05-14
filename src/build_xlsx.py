"""Build the multi-tab outreach-ready Excel file from data/enriched.csv.

Tabs:
    Tier A (DM ready)   - linkedin + valid/catch_all email
    Tier B (Call ready) - phone + valid/catch_all email
    Tier C (Email only) - email only
    Needs Review        - linkedin confidence = needs_review
    All Leads           - everything
    Summary             - counts + yield stats
"""
from __future__ import annotations

import csv
import os
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(os.environ.get("GRIDIRON_DATASET_ROOT") or Path(__file__).resolve().parent.parent)
SRC = ROOT / "data" / "enriched.csv"
OUT_DIR = ROOT / "output"
OUT_DIR.mkdir(exist_ok=True)
OUT = OUT_DIR / "gridiron-leads-enriched.xlsx"

# Order shown in every data tab — marketing-friendly schema
DISPLAY_COLS = [
    "tier", "best_channel",
    "first_name", "last_name", "full_name",
    "role", "role_category",
    "school_name", "school_domain", "city", "state",
    "email",
    "linkedin_url",
    "phone",
    "icp_tier", "icp_score",
    "enriched_at", "source_page", "notes", "lead_id",
]

# Style constants
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TIER_FILLS = {
    "A": PatternFill("solid", fgColor="C6EFCE"),  # green
    "B": PatternFill("solid", fgColor="FFEB9C"),  # yellow
    "C": PatternFill("solid", fgColor="FFE4B5"),  # peach
    "D": PatternFill("solid", fgColor="FCE4D6"),  # orange
    "E": PatternFill("solid", fgColor="F2F2F2"),  # gray
}
EMAIL_FILLS = {
    "valid":     PatternFill("solid", fgColor="C6EFCE"),
    "catch_all": PatternFill("solid", fgColor="FFEB9C"),
    "invalid":   PatternFill("solid", fgColor="FFC7CE"),
    "unknown":   PatternFill("solid", fgColor="E7E6E6"),
}
THIN_BORDER = Border(bottom=Side(style="thin", color="DDDDDD"))


def col_widths_for(rows: list[dict]) -> dict[str, int]:
    widths = {}
    for c in DISPLAY_COLS:
        m = len(c)
        for r in rows[:80]:
            m = max(m, min(len(str(r.get(c, ""))), 60))
        widths[c] = m + 2
    # specific overrides
    widths["lead_id"] = 12  # we're hiding it anyway
    widths["notes"] = 40
    widths["source_page"] = 35
    widths["linkedin_url"] = 38
    widths["school_name"] = 30
    widths["email"] = 32
    return widths


def write_sheet(wb: Workbook, name: str, rows: list[dict], widths: dict[str, int]) -> None:
    ws = wb.create_sheet(name)
    # header
    for col_idx, col in enumerate(DISPLAY_COLS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col.replace("_", " ").title())
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
    # rows
    for row_idx, row in enumerate(rows, 2):
        for col_idx, col in enumerate(DISPLAY_COLS, 1):
            val = row.get(col, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=False)
            cell.border = THIN_BORDER

            if col == "tier" and val in TIER_FILLS:
                cell.fill = TIER_FILLS[val]
                cell.font = Font(bold=True)
            if col == "linkedin_url" and val:
                cell.hyperlink = val
                cell.font = Font(color="0563C1", underline="single")
            if col == "email" and val:
                cell.hyperlink = f"mailto:{val}"
                cell.font = Font(color="0563C1", underline="single")
            if col == "source_page" and val:
                cell.hyperlink = val
                cell.font = Font(color="0563C1", underline="single")
            if col == "phone" and val:
                cell.hyperlink = f"tel:{val}"
                cell.font = Font(color="0563C1", underline="single")

    # widths, freeze, autofilter
    for col_idx, col in enumerate(DISPLAY_COLS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(col, 14)
    ws.freeze_panes = "A2"
    if rows:
        last_col = get_column_letter(len(DISPLAY_COLS))
        ws.auto_filter.ref = f"A1:{last_col}{len(rows) + 1}"
    ws.row_dimensions[1].height = 22

    # Hide noisy columns by default
    for col_idx, col in enumerate(DISPLAY_COLS, 1):
        if col in ("lead_id",):
            ws.column_dimensions[get_column_letter(col_idx)].hidden = True


def write_summary(wb: Workbook, all_rows: list[dict]) -> None:
    ws = wb.create_sheet("Summary", 0)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 32

    title_font = Font(bold=True, size=14, color="1F4E78")
    h2_font = Font(bold=True, size=11, color="1F4E78")

    ws["A1"] = "Gridiron Leads – Enrichment Summary"
    ws["A1"].font = title_font
    ws["A2"] = f"Source: gridiron-leads-2026-05-08.xlsx (synthetic rows excluded)"
    ws["A2"].font = Font(italic=True, color="666666")

    row = 4
    ws.cell(row=row, column=1, value="Total leads (real, deduped)").font = h2_font
    ws.cell(row=row, column=2, value=len(all_rows))
    row += 2

    ws.cell(row=row, column=1, value="By outreach tier").font = h2_font
    row += 1
    tier_labels = {
        "A": "A — LinkedIn + valid email (DM ready)",
        "B": "B — Phone + valid email (Call ready)",
        "C": "C — Email only",
        "D": "D — LinkedIn or phone only, bad email",
        "E": "E — Nothing usable",
    }
    tiers = Counter(r["tier"] for r in all_rows)
    for t in ["A", "B", "C", "D", "E"]:
        ws.cell(row=row, column=1, value=tier_labels[t])
        ws.cell(row=row, column=2, value=tiers.get(t, 0))
        if t in TIER_FILLS:
            ws.cell(row=row, column=1).fill = TIER_FILLS[t]
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="Email verification").font = h2_font
    row += 1
    statuses = Counter(r["email_status"] or "missing" for r in all_rows)
    for s, n in statuses.most_common():
        ws.cell(row=row, column=1, value=s)
        ws.cell(row=row, column=2, value=n)
        if s in EMAIL_FILLS:
            ws.cell(row=row, column=1).fill = EMAIL_FILLS[s]
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="LinkedIn discovery").font = h2_font
    row += 1
    li = Counter(r["linkedin_confidence"] or "no_url_found" for r in all_rows)
    for s in ["high", "medium", "needs_review", "no_url_found"]:
        if li.get(s):
            ws.cell(row=row, column=1, value=s)
            ws.cell(row=row, column=2, value=li[s])
            row += 1
    row += 1

    ws.cell(row=row, column=1, value="Phone coverage").font = h2_font
    row += 1
    school_p = sum(1 for r in all_rows if r["phone_school"])
    direct_p = sum(1 for r in all_rows if r["phone_direct"])
    ws.cell(row=row, column=1, value="With school/athletic dept phone")
    ws.cell(row=row, column=2, value=school_p)
    row += 1
    ws.cell(row=row, column=1, value="With coach-direct phone")
    ws.cell(row=row, column=2, value=direct_p)
    row += 1
    ws.cell(row=row, column=1, value="With any phone")
    ws.cell(row=row, column=2, value=sum(1 for r in all_rows if r["phone_school"] or r["phone_direct"]))
    row += 2

    ws.cell(row=row, column=1, value="Top 10 states by lead count").font = h2_font
    row += 1
    states = Counter(r["state"] for r in all_rows if r["state"]).most_common(10)
    for st, n in states:
        ws.cell(row=row, column=1, value=st)
        ws.cell(row=row, column=2, value=n)
        row += 1

    ws.sheet_view.showGridLines = False


def main() -> None:
    with SRC.open() as f:
        rows = list(csv.DictReader(f))
    if not rows:
        raise SystemExit("no rows in enriched.csv")

    # Collapse phone_direct + phone_school into a single 'phone' column
    # (preferring the direct line when present)
    for r in rows:
        r["phone"] = (r.get("phone_direct") or "").strip() or (r.get("phone_school") or "").strip()

    widths = col_widths_for(rows)

    tier_a = [r for r in rows if r["tier"] == "A"]
    tier_b = [r for r in rows if r["tier"] == "B"]
    tier_c = [r for r in rows if r["tier"] == "C"]
    needs_review = [r for r in rows if r["linkedin_confidence"] == "needs_review" and r["linkedin_url"]]

    wb = Workbook()
    wb.remove(wb.active)  # drop default sheet

    write_summary(wb, rows)
    write_sheet(wb, "Tier A (DM ready)", tier_a, widths)
    write_sheet(wb, "Tier B (Call ready)", tier_b, widths)
    write_sheet(wb, "Tier C (Email only)", tier_c, widths)
    write_sheet(wb, "Needs Review", needs_review, widths)
    write_sheet(wb, "All Leads", rows, widths)

    wb.save(OUT)
    print(f"wrote {OUT} ({len(rows)} leads | A={len(tier_a)} B={len(tier_b)} C={len(tier_c)} review={len(needs_review)})")


if __name__ == "__main__":
    main()
