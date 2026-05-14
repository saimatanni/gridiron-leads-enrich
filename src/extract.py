"""Extract clean outreach-ready lead schema from the source workbook.

Schema-tolerant: auto-detects common column-name variants for each canonical
field. If the source already has a value for a field (e.g. linkedin_url or
phone), pre-populates the corresponding results file so the enrichment phase
skips that lead and we keep the user's value.
"""
from __future__ import annotations

import csv
import hashlib
import os
import re
from pathlib import Path

import openpyxl

ROOT = Path(os.environ.get("GRIDIRON_DATASET_ROOT") or Path(__file__).resolve().parent.parent)
SRC = ROOT / "source.xlsx"
DATA = ROOT / "data"
DATA.mkdir(exist_ok=True)
OUT = DATA / "leads_clean.csv"

SCHEMA = [
    "lead_id", "tier", "first_name", "last_name", "full_name",
    "role", "role_category", "school_name", "school_domain",
    "city", "state",
    "email", "email_status",
    "linkedin_url", "linkedin_confidence",
    "phone_school", "phone_direct",
    "best_channel", "icp_tier", "icp_score",
    "enriched_at", "source_page", "notes",
]

# Aliases for each canonical field. Match is case-insensitive on the entire
# trimmed header string. First match wins.
ALIASES: dict[str, list[str]] = {
    "id":             ["id", "lead_id", "_id", "identifier", "uuid"],
    "source":         ["source", "source_dataset", "dataset", "origin"],
    "first_name":     ["first_name", "firstname", "first", "given_name", "fname", "first name"],
    "last_name":      ["last_name", "lastname", "last", "family_name", "surname", "lname", "last name"],
    "full_name":      ["full_name", "name", "fullname", "full name"],
    "email":          ["email", "e-mail", "email_address", "email address"],
    "title":          ["title", "designation", "role", "position", "job_title", "job title"],
    "company":        ["company", "school_name", "school", "institute", "institution",
                       "organization", "org", "school name", "company_name"],
    "company_domain": ["company_domain", "domain", "school_domain", "website"],
    "state":          ["state", "region"],
    "city":           ["city", "town"],
    "location":       ["location", "address", "city_state"],
    "_phone":         ["_phone", "phone", "phone_number", "telephone", "contact_number"],
    "_phone2":        ["phone2"],  # secondary phone column if present
    "icp_score":      ["icp_match_score", "icp_score", "score"],
    "icp_tier":       ["icp_tier", "tier_source"],
    "linkedin_url":   ["linkedin_url", "linkedin", "linkedin url", "li_url"],
    "contact_type":   ["contact_type", "role_category", "type", "role_type"],
    "_maxpreps_url":  ["_maxpreps_url", "maxpreps_url"],
}

# Words that mark a row as synthetic/fake when present in the `source` column
SYN_MARKERS = ("synthetic", "fake", "test_data", "generated")


def build_column_map(header: list[str]) -> dict[str, int]:
    """Return canonical_field -> column_index (or absent if not found)."""
    lc_index = {(h or "").strip().lower(): i for i, h in enumerate(header)}
    out: dict[str, int] = {}
    for canonical, aliases in ALIASES.items():
        for a in aliases:
            i = lc_index.get(a.lower())
            if i is not None:
                out[canonical] = i
                break
    return out


def parse_location(loc: str) -> tuple[str, str]:
    if not loc:
        return "", ""
    m = re.match(r"^\s*(.+?),\s*([A-Z]{2})\s*$", loc.strip())
    if m:
        return m.group(1).strip(), m.group(2)
    return loc.strip(), ""


def generate_lead_id(email: str, first: str, last: str) -> str:
    """Deterministic lead_id when source doesn't have one — based on email."""
    h = hashlib.md5(f"{email}|{first}|{last}".encode()).hexdigest()
    return f"lead_{h[:24]}"


def main() -> None:
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(rows)]
    col = build_column_map(header)
    print(f"Detected source columns: {len(header)} total")
    print(f"  Mapped to canonical: {sorted(col.keys())}")
    missing = [k for k in ("email", "first_name", "last_name") if k not in col]
    if missing:
        print(f"  ⚠ missing required columns: {missing}")

    def g(row: tuple, canonical: str) -> str:
        i = col.get(canonical)
        if i is None or i >= len(row):
            return ""
        v = row[i]
        return "" if v is None else str(v).strip()

    cleaned: list[dict] = []
    seen: set[tuple[str, str, str]] = set()
    real_count = 0
    drop_dup = 0
    drop_no_email = 0
    syn_dropped = 0

    # Pre-populated enrichment outputs (so downstream scripts skip leads where
    # we already have a value from the source file)
    pre_linkedin: list[dict] = []
    pre_phone: list[dict] = []

    # Synthetic rows captured so the Fake-users tab can show them right away
    # (we no longer have to wait for the final add_synthetic_tabs step)
    syn_rows: list[dict] = []

    for row in rows:
        if not row:
            continue

        source_val = g(row, "source").lower()
        if any(m in source_val for m in SYN_MARKERS):
            syn_dropped += 1
            # Capture for the Fake-users tab so it can populate immediately
            syn_rows.append({
                "dup_count": "",  # filled below
                "email": g(row, "email").lower(),
                "first_name": g(row, "first_name"),
                "last_name": g(row, "last_name"),
                "company": g(row, "company"),
                "state": g(row, "state"),
                "enrichment_status": "",
                "source": source_val,
                "company_domain": g(row, "company_domain"),
                "id": g(row, "id"),
            })
            continue
        real_count += 1

        first = g(row, "first_name")
        last = g(row, "last_name")
        full = g(row, "full_name")
        if not first and not last and full:
            # Try splitting "Scott Wooster" into first/last
            parts = full.split(maxsplit=1)
            first = parts[0]
            last = parts[1] if len(parts) > 1 else ""

        email = g(row, "email").lower()
        if not email:
            drop_no_email += 1
            continue
        school = g(row, "company")
        key = (school.lower(), first.lower(), last.lower())
        if key in seen:
            drop_dup += 1
            continue
        seen.add(key)

        # Generate lead_id if source doesn't have one
        lead_id = g(row, "id") or generate_lead_id(email, first, last)

        city_loc, state_loc = parse_location(g(row, "location"))
        state = g(row, "state") or state_loc
        city = g(row, "city") or city_loc

        existing_linkedin = g(row, "linkedin_url")
        existing_phone = g(row, "_phone") or g(row, "_phone2")

        cleaned.append({
            "lead_id": lead_id,
            "tier": "",
            "first_name": first,
            "last_name": last,
            "full_name": full or f"{first} {last}".strip(),
            "role": g(row, "title"),
            "role_category": g(row, "contact_type"),
            "school_name": school,
            "school_domain": g(row, "company_domain"),
            "city": city,
            "state": state,
            "email": email,
            "email_status": "",
            "linkedin_url": existing_linkedin,
            "linkedin_confidence": "from_source" if existing_linkedin else "",
            "phone_school": existing_phone,
            "phone_direct": "",
            "best_channel": "",
            "icp_tier": g(row, "icp_tier"),
            "icp_score": g(row, "icp_score"),
            "enriched_at": "",
            "source_page": "",
            "notes": "",
        })

        # Pre-populate the LinkedIn result file so linkedin_find skips this lead
        if existing_linkedin:
            pre_linkedin.append({
                "lead_id": lead_id,
                "linkedin_url": existing_linkedin,
                "linkedin_confidence": "from_source",
                "source_query": "(from source file)",
                "snippet": "",
                "engine": "",
            })

        # Pre-populate phone result so phone_find skips this lead
        if existing_phone:
            pre_phone.append({
                "lead_id": lead_id,
                "phone_school": existing_phone,
                "phone_direct": "",
                "source_page": "(from source file)",
                "notes": "phone from source",
            })

    with OUT.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=SCHEMA)
        w.writeheader()
        w.writerows(cleaned)

    if pre_linkedin:
        li_path = DATA / "linkedin_results.csv"
        with li_path.open("w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=["lead_id", "linkedin_url", "linkedin_confidence",
                                              "source_query", "snippet", "engine"])
            w.writeheader()
            w.writerows(pre_linkedin)
        print(f"  pre-populated linkedin_results.csv with {len(pre_linkedin)} rows from source")

    if pre_phone:
        ph_path = DATA / "phone_results.csv"
        with ph_path.open("w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=["lead_id", "phone_school", "phone_direct",
                                              "source_page", "notes"])
            w.writeheader()
            w.writerows(pre_phone)
        print(f"  pre-populated phone_results.csv with {len(pre_phone)} rows from source")

    # Dump synthetic rows immediately so the Fake-users tab populates from
    # the very first phase, not after the final add_synthetic step
    if syn_rows:
        import collections as _c
        out_dir = ROOT / "output"
        out_dir.mkdir(exist_ok=True)
        # Compute dup_count per email
        counts = _c.Counter(r["email"] for r in syn_rows if r["email"])
        for r in syn_rows:
            r["dup_count"] = counts.get(r["email"], 0)
        # Sort by duplicate count desc so worst offenders are at top
        syn_rows.sort(key=lambda r: -r["dup_count"])
        syn_csv = out_dir / "synthetic-fake-users.csv"
        with syn_csv.open("w", newline="") as f:
            cols = ["dup_count", "email", "first_name", "last_name", "company",
                    "state", "enrichment_status", "source", "company_domain", "id"]
            w = csv.DictWriter(f, fieldnames=cols)
            w.writeheader()
            w.writerows(syn_rows)
        print(f"  wrote {syn_csv} with {len(syn_rows)} synthetic rows (Fake-users tab will populate immediately)")

    print(f"source rows total: {real_count + syn_dropped}")
    print(f"  flagged synthetic/fake (dropped): {syn_dropped}")
    print(f"  real rows: {real_count}")
    print(f"  dropped (no email): {drop_no_email}")
    print(f"  dropped (duplicate school+name): {drop_dup}")
    print(f"  kept: {len(cleaned)}")
    print(f"  with school domain: {sum(1 for r in cleaned if r['school_domain'])}")
    print(f"  with phone from source: {sum(1 for r in cleaned if r['phone_school'])}")
    print(f"  with linkedin from source: {sum(1 for r in cleaned if r['linkedin_url'])}")
    print(f"written: {OUT}")


if __name__ == "__main__":
    main()
